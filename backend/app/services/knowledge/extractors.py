"""Extracts plain text from each supported knowledge source type.

Each function takes raw bytes (or, for URLs, fetched HTML) and returns
plain text ready for chunking. Kept separate from KnowledgeService so
each extractor is independently unit-testable and independently
replaceable (e.g. swapping in OCR for scanned PDFs later touches only
extract_text_from_pdf).
"""
import io

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

from app.models.knowledge import KnowledgeSourceType


class ExtractionError(Exception):
    pass


def extract_text_from_plain_bytes(raw: bytes) -> str:
    """Used for TEXT, MARKDOWN, and FAQ source types — all are just
    UTF-8 text, markdown syntax and Q&A formatting are left as-is since
    the AI provider handles that structure fine at generation time."""
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("utf-8", errors="replace")


def extract_text_from_pdf(raw: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(raw))
    except Exception as exc:
        raise ExtractionError(f"Could not open PDF: {exc}") from exc

    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)

    if not pages:
        raise ExtractionError("No extractable text found in PDF (it may be a scanned image without OCR).")
    return "\n\n".join(pages)


def extract_text_from_docx(raw: bytes) -> str:
    try:
        document = DocxDocument(io.BytesIO(raw))
    except Exception as exc:
        raise ExtractionError(f"Could not open Word document: {exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    if not parts:
        raise ExtractionError("No extractable text found in the Word document.")
    return "\n\n".join(parts)


def extract_text_from_html(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        raise ExtractionError("No extractable text found on the page.")
    return "\n".join(lines)


_EXTRACTORS_BY_SOURCE_TYPE = {
    KnowledgeSourceType.TEXT: extract_text_from_plain_bytes,
    KnowledgeSourceType.MARKDOWN: extract_text_from_plain_bytes,
    KnowledgeSourceType.FAQ: extract_text_from_plain_bytes,
    KnowledgeSourceType.PDF: extract_text_from_pdf,
    KnowledgeSourceType.DOCX: extract_text_from_docx,
}


def extract_text(source_type: KnowledgeSourceType, raw: bytes) -> str:
    """Dispatch for file-based sources. URL sources go through
    extract_text_from_html directly since they start as fetched HTML,
    not raw bytes from an upload."""
    extractor = _EXTRACTORS_BY_SOURCE_TYPE.get(source_type)
    if extractor is None:
        raise ExtractionError(f"No extractor available for source type '{source_type.value}'")
    return extractor(raw)


def infer_source_type_from_filename(filename: str) -> KnowledgeSourceType:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return KnowledgeSourceType.PDF
    if lower.endswith(".docx"):
        return KnowledgeSourceType.DOCX
    if lower.endswith(".md") or lower.endswith(".markdown"):
        return KnowledgeSourceType.MARKDOWN
    return KnowledgeSourceType.TEXT
